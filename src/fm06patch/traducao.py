"""Textos do jogo em português do Brasil (glossário de data/modelos/traducao_ptbr.xlsx).

Dois arquivos, nos dois o texto novo tem de caber no antigo (é completado com espaços):

- lang_db.dat: uma seção por idioma; no começo do arquivo, a lista (u32 idioma, u32
  posição) a partir de 0x12; o português de Portugal é o idioma 0x13
  (0x014AB8CD..0x01795075). Cada texto é u32 tamanho (em caracteres), UTF-16LE, 00 00.
  Na seção ficam nomes de clubes e cidades e os textos traduzidos (prêmios, lesões,
  fases, descrições de jornal). Só mudam os textos que não aparecem igual em nenhum
  outro idioma: nome próprio ("Presidente Prudente (SP)") é igual em todos, texto
  traduzido não. O arquivo tem posições absolutas das seções, por isso não pode crescer.
- data/languages/portuguese.ltc: as telas (ltc.py). Lá só há texto, sem nomes próprios.

Cada arquivo tem suas regras (Rules): termos do glossário que não servem ali, trocas a
mais e as palavras que mudam de gênero ("equipa" -> "time" muda o artigo antes: "a
equipa" -> "o time"; com adjetivo concordando depois, o texto não muda).
"""

from __future__ import annotations

import re
import struct
from dataclasses import dataclass, field

PORTUGUESE = 0x13
_TEXT = re.compile(rb"([\x01-\xff])\x00\x00\x00((?:[^\x00]\x00|[\x00-\xff][\x01-\xff])+?)\x00\x00", re.S)
_TAG = re.compile(r"<[^>]*>|\{[^}]*\}|\[[^\]]*\]")       # <%team#1>, {do}, [COMMENT ...]

# grafia de Portugal -> do Brasil (mesmo gênero; quase sempre fica mais curto)
ORTHOGRAPHY = {
    "actual": "atual", "actuais": "atuais", "actualmente": "atualmente", "actualizar": "atualizar",
    "actualização": "atualização", "actualizado": "atualizado", "actualizada": "atualizada",
    "acção": "ação", "acções": "ações", "activo": "ativo", "activa": "ativa", "activos": "ativos",
    "activas": "ativas", "actividade": "atividade", "actividades": "atividades",
    "objectivo": "objetivo", "objectivos": "objetivos", "director": "diretor", "directores": "diretores",
    "directora": "diretora", "directamente": "diretamente", "directo": "direto", "directa": "direta",
    "excepção": "exceção", "excepções": "exceções", "excepto": "exceto", "exacto": "exato",
    "exacta": "exata", "exactamente": "exatamente", "óptimo": "ótimo", "óptima": "ótima",
    "óptimos": "ótimos", "facto": "fato", "contacto": "contato", "contactos": "contatos",
    "contactar": "contatar", "projecto": "projeto", "projectos": "projetos", "sector": "setor",
    "sectores": "setores", "reacção": "reação", "protecção": "proteção", "eléctrico": "elétrico",
    "colecção": "coleção", "correcto": "correto", "correcta": "correta", "incorrecto": "incorreto",
    "efectuar": "efetuar", "efectuado": "efetuado", "efectuada": "efetuada", "efectivo": "efetivo",
    "adoptar": "adotar", "registo": "registro", "registar": "registrar", "registado": "registrado",
    "económico": "econômico", "económica": "econômica", "económicos": "econômicos",
    "económicas": "econômicas", "fenómeno": "fenômeno", "génio": "gênio", "anónimo": "anônimo",
    "electrónico": "eletrônico", "electrónica": "eletrônica", "polémica": "polêmica",
    "polémico": "polêmico", "polémicas": "polêmicas", "direcção": "direção", "direcções": "direções",
    "seleccionar": "selecionar", "seleccionado": "selecionado", "seleccionada": "selecionada",
    "seleccionados": "selecionados", "seleccionadas": "selecionadas", "seleccionador": "selecionador",
    "correctos": "corretos", "correctas": "corretas", "correctamente": "corretamente", "directos": "diretos",
    "directas": "diretas", "actuar": "atuar", "actuação": "atuação", "actuações": "atuações",
    "exactos": "exatos", "óptimas": "ótimas", "efectivamente": "efetivamente",
}

_TO_M = {"a": "o", "as": "os", "uma": "um", "umas": "uns", "da": "do", "das": "dos", "na": "no", "nas": "nos",
         "pela": "pelo", "pelas": "pelos", "à": "ao", "às": "aos", "esta": "este", "estas": "estes",
         "essa": "esse", "essas": "esses", "sua": "seu", "suas": "seus", "nova": "novo", "novas": "novos",
         "antiga": "antigo", "primeira": "primeiro", "última": "último", "outra": "outro", "outras": "outros",
         "nenhuma": "nenhum", "toda": "todo", "todas": "todos", "minha": "meu", "nossa": "nosso",
         "alguma": "algum", "algumas": "alguns"}
_TO_F = {v: k for k, v in _TO_M.items()}
_FUNCTION = {"a", "à", "as", "às", "o", "os", "ao", "aos", "da", "das", "do", "dos", "na", "nas", "no", "nos",
             "para", "pela", "pelas", "pelo", "pelos", "contra", "sobre", "como", "ou", "e", "que", "se", "sem",
             "com", "de", "em", "por", "entre", "até", "após", "durante"}


@dataclass
class Rules:
    skip: set[str] = field(default_factory=set)          # termos do glossário (minúsculos) que não servem
    extra: dict[str, str] = field(default_factory=dict)  # trocas a mais (ganham do glossário)
    gender: dict[str, str] = field(default_factory=dict)  # termo -> gênero novo ("m"/"f") quando muda


LANG_DB = Rules(
    skip={"partida",              # "perna partida" é perna quebrada
          "mão",                  # "mão partida"; "Primeira Mão" fica com o nomes_fases
          "marcador",             # "Melhor Marcador" é artilheiro (extra), não placar
          "campeonato nacional",   # vale para todos os países
          "subida", "descida"},    # "subida na classificação" não é acesso
    extra={"melhor marcador": "artilheiro", "melhores marcadores": "artilheiros",
           "eliminatória": "fase", "eliminatórias": "fases", "dos adeptos": "da torcida", **ORTHOGRAPHY},
    gender={"equipa": "m", "equipas": "m"},
)

# nas telas, "equipa" vira "equipe" (feminino, como no original): "a equipa principal" ->
# "a equipe principal" sem mexer no resto da frase
SCREENS = Rules(
    skip={"partida", "mão", "marcador", "campeonato nacional", "juniores", "formação", "classificação",
          "vencedor", "encontro", "compensação", "nulo", "seguinte", "correio", "equipa principal",
          "equipa de reservas", "equipa sub-20", "camadas jovens",
          "médio", "médios",          # também é "médio" de média ("Salário Médio"): só as posições (extra)
          "subida", "descida",        # "subida na classificação", "subida do moral": só "de divisão" (extra)
          "desarme"},                 # no Brasil também é desarme
    extra={"equipa": "equipe", "equipas": "equipes", "primeira mão": "ida", "segunda mão": "volta",
           "eliminatória": "fase", "eliminatórias": "fases", "autocarro": "ônibus", "remunerações": "salários",
           "cabeceamento": "cabeceio", "cabeceamentos": "cabeceios",
           "médio ofensivo": "meia atacante", "médios ofensivos": "meias atacantes",
           "médio defensivo": "volante", "médios defensivos": "volantes", "médio direito": "meia direita",
           "médio esquerdo": "meia esquerda", "médio centro": "meia central", "médio (centro)": "meia (centro)",
           "subida de divisão": "acesso", "descida de divisão": "rebaixamento", **ORTHOGRAPHY},
    gender={"grande penalidade": "m", "grandes penalidades": "m", "livre directo": "f", "livre indirecto": "f",
            "pontapé livre": "f", "pontapés livres": "f", "prolongamento": "f", "desempate por penáltis": "f",
            "finta": "m", "fintas": "m", "passe em profundidade": "f", "castigo": "f",
            "subida de divisão": "m", "descida de divisão": "m", "directório": "f", "ecrã": "f",
            "médio direito": "f", "médio esquerdo": "f", "médio centro": "f"},
)


@dataclass
class Change:
    at: int          # posição do texto (depois do u32 tamanho)
    old: str
    new: str


def sections(data: bytes) -> dict[int, tuple[int, int]]:
    """idioma -> (início, fim) da seção do lang_db."""
    found, at = [], 0x12
    while True:
        language, start = struct.unpack_from("<II", data, at)
        if not 0x50 < start < len(data) or (found and start <= found[-1][1]):
            break
        found.append((language, start))
        at += 8
    ends = [start for _, start in found[1:]] + [len(data)]
    return {language: (start, end) for (language, start), end in zip(found, ends)}


def texts(data: bytes, start: int, end: int) -> list[tuple[int, str]]:
    out = []
    for match in _TEXT.finditer(data, start, end):
        body = match.group(2)
        if len(body) != 2 * match.group(1)[0]:
            continue
        try:
            text = body.decode("utf-16-le")
        except UnicodeDecodeError:
            continue
        if text.isprintable():
            out.append((match.start(2), text))
    return out


_SMALL = {"a", "o", "as", "os", "e", "de", "da", "do", "das", "dos", "em", "na", "no", "por", "para", "com"}


def _cased(source: str, target: str) -> str:
    if source.isupper() and len(source) > 1:
        return target.upper()
    words = source.split()
    if len(words) > 1 and all(w[:1].isupper() for w in words if w.casefold() not in _SMALL):   # "Carregar Jogo"
        return " ".join(w if w.casefold() in _SMALL and i else w[:1].upper() + w[1:]
                        for i, w in enumerate(target.split()))
    if source[:1].isupper():
        return target[:1].upper() + target[1:]
    if source[:1].islower() and not target[1:2].isupper():      # "época" com a alternativa "Época"
        return target[:1].lower() + target[1:]
    return target


def _ends_like(word: str, gender: str) -> bool:
    """A palavra concorda com o gênero (termina em -a/-as ou -o/-os)?"""
    word = word.casefold()
    return len(word) > 2 and word not in _FUNCTION and word.endswith(("a", "as") if gender == "f" else ("o", "os"))


class Glossary:
    """pt_pt -> pt_br (palavras e expressões, as mais longas primeiro, palavra inteira, sem
    distinguir maiúsculas: a maiúscula do texto vai para a tradução)."""

    def __init__(self, terms: dict[str, str], short: dict[str, str] | None = None, rules: Rules = LANG_DB):
        self.rules = rules
        overridden = {k.casefold() for k in rules.extra}
        self.terms = {k: v for k, v in terms.items()
                      if k.casefold() not in rules.skip and k.casefold() not in overridden and k != v}
        self.terms.update(rules.extra)
        self.short = {k.casefold(): v for k, v in (short or {}).items()}
        self.folded = {}
        for key, value in self.terms.items():
            self.folded.setdefault(key.casefold(), value)
        keys = sorted(self.terms, key=len, reverse=True)
        self.pattern = re.compile(r"(?<!\w)(" + "|".join(re.escape(k) for k in keys) + r")(?!\w)", re.I)

    def _value(self, found: str, short: bool) -> str:
        if short and found.casefold() in self.short:
            return _cased(found, self.short[found.casefold()])
        value = self.terms[found] if found in self.terms else _cased(found, self.folded[found.casefold()])
        if short and len(value) > len(found):       # sem alternativa curta: o termo fica como está
            return found
        return value

    def translate(self, text: str, short: bool = False) -> str | None:
        """Texto traduzido, ou None se não dá para traduzir com segurança."""
        tags = [m.span() for m in _TAG.finditer(text)]
        pieces, last = [], 0
        for match in self.pattern.finditer(text):
            if any(a < match.end() and match.start() < b for a, b in tags):
                continue
            found = match.group(1)
            new = self._value(found, short)
            before = text[last:match.start()]
            gender = self.rules.gender.get(found.casefold())
            if gender and new.casefold() != found.casefold():   # a alternativa curta pode ser o próprio termo
                old = "f" if gender == "m" else "m"
                after = text[match.end():]
                if after.lstrip().startswith("<%"):                 # adjetivo de nacionalidade concorda
                    return None
                following = re.match(r"\s+(\w+)", after)
                if following and _ends_like(following.group(1), old):
                    return None
                head = re.search(r"(\w+)(\s+)$", before)
                if head:
                    word = head.group(1)
                    table = _TO_M if gender == "m" else _TO_F
                    if word.casefold() in table:
                        before = before[:head.start()] + _cased(word, table[word.casefold()]) + head.group(2)
                    elif _ends_like(word, old):
                        return None
            pieces += [before, new]
            last = match.end()
        if not pieces:
            return text
        return "".join(pieces) + text[last:]

    def fit(self, old: str, phrases: dict[str, str] | None = None,
            short_phrases: dict[str, str] | None = None) -> tuple[str | None, bool]:
        """(texto novo, coube?) para um texto do jogo; None = não muda."""
        candidates = [(phrases or {}).get(old), (short_phrases or {}).get(old)]
        candidates += [self.translate(old), self.translate(old, short=True)]
        tried = None
        for new in candidates:
            if new is None or new == old:
                continue
            if len(new) <= len(old):
                return new, True
            tried = tried or new
        return tried, False


def plan(data: bytes, glossary: Glossary, phrases: dict[str, str] | None = None,
         short_phrases: dict[str, str] | None = None) -> tuple[list[Change], list[str]]:
    """lang_db: textos a trocar e os que não couberam."""
    bounds = sections(data)
    others: set[str] = set()
    for language, (start, end) in bounds.items():
        if language != PORTUGUESE:
            others.update(text for _, text in texts(data, start, end))
    changes, too_long = [], []
    for at, old in texts(data, *bounds[PORTUGUESE]):
        if old in others:
            continue
        new, fits = glossary.fit(old, phrases, short_phrases)
        if new is None:
            continue
        if fits:
            changes.append(Change(at, old, new))
        else:
            too_long.append(f"{old} -> {new}")
    return changes, too_long


def plan_screens(data: bytes, glossary: Glossary, phrases: dict[str, str] | None = None,
                 short_phrases: dict[str, str] | None = None) -> tuple[list[Change], list[str]]:
    """portuguese.ltc: traduções a trocar e as que não couberam."""
    from .ltc import translations

    changes, too_long = [], []
    for at, old in translations(data):
        new, fits = glossary.fit(old, phrases, short_phrases)
        if new is None:
            continue
        if fits:
            changes.append(Change(at, old, new))
        else:
            too_long.append(f"{old} -> {new}")
    return changes, too_long


def apply(data: bytes, changes: list[Change]) -> bytes:
    out = bytearray(data)
    for change in changes:
        out[change.at:change.at + 2 * len(change.old)] = change.new.ljust(len(change.old)).encode("utf-16-le")
    return bytes(out)


@dataclass
class Sheet:
    terms: dict[str, str]
    short: dict[str, str]
    phrases: dict[str, str]
    short_phrases: dict[str, str]

    def glossary(self, rules: Rules = LANG_DB) -> Glossary:
        return Glossary(self.terms, self.short, rules)


def load_sheet(path) -> Sheet:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True)
    terms, short = {}, {}
    for row in book["Glossario"].iter_rows(min_row=2, values_only=True):
        pt_pt, pt_br = str(row[0] or "").strip(), str(row[1] or "").strip()
        if pt_pt and pt_br:
            terms[pt_pt] = pt_br
            if row[4] and str(row[4]).strip():
                short[pt_pt] = str(row[4]).strip()
    phrases, short_phrases = {}, {}
    sheet = book["Frases"]
    header = [str(c or "") for c in next(sheet.iter_rows(max_row=1, values_only=True))]
    alt = header.index("alternativa_curta") if "alternativa_curta" in header else None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        old, new = str(row[1] or "").strip(), str(row[2] or "").strip()
        if old and new:
            phrases[old] = new
            if alt is not None and row[alt] and str(row[alt]).strip():
                short_phrases[old] = str(row[alt]).strip()
    return Sheet(terms, short, phrases, short_phrases)
