"""Clubes da CONMEBOL de 2026 por planilha (sem o Sofascore).

    python scripts\\estrangeiros_planilha.py --modelo   (cria data\\modelos\\estrangeiros_2026.xlsx)
    python scripts\\estrangeiros_planilha.py            (monta data\\estrangeiros a partir dela)

A planilha tem uma aba Clubes, com todos os clubes do FM06 de Argentina, Uruguai,
Colômbia, Paraguai, Equador, Chile, Peru, Bolívia e Venezuela (os que acho que estão na
primeira divisão de 2026 já vêm com "sim" em usar), uma aba Jogadores (o elenco de 2026
de cada clube marcado) e uma aba Carreiras (opcional). As instruções estão na aba LEIA.

Grava em data\\estrangeiros os mesmos arquivos que o coletar_estrangeiros.py gravaria
(identidade, jogadores, técnicos e carreiras), que o aplicar_tudo.py usa. O CA e o PA
vazios saem do valor de mercado, da idade e da "série" do país, como nos brasileiros
(montar_jogadores.py); peso, altura, pé e salário vazios são estimados.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from fm06patch import estrangeiros as ex, people_contracts  # noqa: E402
from fm06patch.squad_builder import ROLE_POSITIONS  # noqa: E402
from montar_carreiras import FIELDS as CAREER_FIELDS  # noqa: E402
from montar_estrangeiros import COACH_FIELDS, server_bytes  # noqa: E402
import montar_jogadores as mj  # noqa: E402

MODEL = ROOT / "data" / "modelos" / "estrangeiros_2026.xlsx"
PEOPLE_BACKUP = ROOT / "workspace" / "baseline" / "people_db.original.dat"
BY_CODE = {v[1]: k for k, v in ex.NATIONS.items()}       # "ARG" -> "AR"
CLUB_COLUMNS = ["pais", "clube_fm06", "nome_curto_fm06", "uid_jogo", "jogadores_no_fm06", "usar", "clube_2026",
                "tecnico", "tecnico_nacionalidade", "tecnico_nascimento", "fonte", "obs"]
PLAYER_COLUMNS = ["clube_2026", "numero", "nome_comum", "primeiro_nome", "sobrenome", "nascimento", "nacionalidade",
                  "posicoes", "pe", "altura", "valor_mercado_eur", "ca", "pa", "contrato_fim", "jogos_selecao",
                  "gols_selecao", "fonte", "obs"]
CAREER_COLUMNS = ["clube_2026", "nome_comum", "nascimento", "temporada", "clube", "emprestimo", "valor_eur", "fonte"]
# clubes que acho que estão na primeira divisão de 2026 (nome do FM06 -> nome de 2026); confira
FIRST_2026 = {
    "AR": {"River Plate": "River Plate", "Boca Juniors": "Boca Juniors", "Racing Club": "Racing Club",
           "Independiente": "Independiente", "San Lorenzo de Almagro": "San Lorenzo",
           "Estudiantes de La Plata": "Estudiantes", "Gimnasia y Esgrima de La Plata": "Gimnasia LP",
           "Vélez Sarsfield": "Vélez Sarsfield", "Talleres de Córdoba": "Talleres", "Belgrano de Córdoba": "Belgrano",
           "Rosario Central": "Rosario Central", "Newell's Old Boys": "Newell's Old Boys", "Lanús": "Lanús",
           "Banfield": "Banfield", "Huracán": "Huracán", "Argentinos Juniors": "Argentinos Juniors",
           "Platense": "Platense", "Defensa y Justicia": "Defensa y Justicia", "Godoy Cruz Antonio Tomba": "Godoy Cruz",
           "Unión de Santa Fe": "Unión", "Atlético Tucumán": "Atlético Tucumán", "Instituto de Córdoba": "Instituto",
           "Independiente Rivadavia de Mendoza": "Independiente Rivadavia", "Barracas Central": "Barracas Central",
           "Deportivo Riestra": "Deportivo Riestra", "San Martín de San Juan": "San Martín (SJ)",
           "Aldosivi": "Aldosivi", "Gimnasia y Esgrima de Mendoza": "Gimnasia Mendoza"},
    "UY": {"CA Peñarol": "Peñarol", "C Nacional F": "Nacional", "Defensor Sporting C": "Defensor Sporting",
           "Liverpool FC": "Liverpool", "Danubio FC": "Danubio", "Montevideo Wanderers FC": "Wanderers",
           "CA Cerro": "Cerro", "Racing C Montevideo": "Racing", "CA Progreso": "Progreso",
           "Cerro Largo FC": "Cerro Largo", "CSD Juventud de Las Piedras": "Juventud", "CA River Plate": "River Plate",
           "Central Español FC": "Central Español"},
    "CO": {"CDC Atlético Nacional": "Atlético Nacional", "CD Los Millonarios": "Millonarios",
           "CD América de Cali": "América de Cali", "AC Deportivo Cali": "Deportivo Cali", "CPD Junior": "Junior",
           "Independiente Santa Fé": "Santa Fe", "CDI Medellín": "Independiente Medellín", "CD Once Caldas": "Once Caldas",
           "CCD Tolima": "Deportes Tolima", "CA Bucaramanga": "Atlético Bucaramanga",
           "Deportivo Pereira FC": "Deportivo Pereira", "Envigado FC": "Envigado", "Alianza Petrolera AS": "Alianza FC",
           "CD La Equidad S": "La Equidad", "CN Cúcuta Deportivo": "Cúcuta Deportivo"},
    "PY": {"C Olimpia": "Olimpia", "C Cerro Porteño": "Cerro Porteño", "C Libertad": "Libertad", "C Guaraní": "Guaraní",
           "C Nacional FC": "Nacional", "C Sportivo Luqueño": "Sportivo Luqueño",
           "C Sportivo Trinidense": "Sportivo Trinidense", "C 2 de Mayo": "2 de Mayo",
           "C Sportivo Ameliano": "Sportivo Ameliano", "Tacuary FBC": "Tacuary", "C Deportivo Recoleta": "Recoleta",
           "C Rubio Ñu": "Rubio Ñu"},
    "EC": {"LDU de Quito": "LDU Quito", "Barcelona SC": "Barcelona SC", "CS Emelec": "Emelec", "SD Aucas": "Aucas",
           "CS Deportivo Cuenca": "Deportivo Cuenca", "CD El Nacional": "El Nacional", "Delfín SC": "Delfín",
           "CSD Macará": "Macará", "CD Técnico Universitario": "Técnico Universitario",
           "CDU Católica": "Universidad Católica"},
    "CL": {"CSD Colo Colo": "Colo-Colo", "CF Universidad de Chile": "Universidad de Chile",
           "CD Universidad Católica": "Universidad Católica", "CD Cobreloa": "Cobreloa",
           "O'Higgins de Rancagua": "O'Higgins", "CD Huachipato": "Huachipato", "CD Palestino": "Palestino",
           "Audax CS Italiano": "Audax Italiano", "CD Unión Española": "Unión Española",
           "CD Everton de Viña del Mar": "Everton", "CD Coquimbo Unido": "Coquimbo Unido",
           "CD Municipal Iquique": "Deportes Iquique", "CD Ñublense SA": "Ñublense", "CD La Serena": "Deportes La Serena",
           "CCD Universidad de Concepción": "Universidad de Concepción", "CD Cobresal": "Cobresal"},
    "PE": {"C Alianza Lima": "Alianza Lima", "C Universitario D": "Universitario", "C Sporting Cristal": "Sporting Cristal",
           "CFBC Melgar": "Melgar", "ADFPC Cienciano": "Cienciano", "C Alianza Atlético": "Alianza Atlético",
           "C Atlético Grau": "Atlético Grau", "Deportivo Garcilaso": "Deportivo Garcilaso", "C Sport Boys A": "Sport Boys",
           "U Técnica de Cajamarca": "UTC"},
    "BO": {"C Bolívar": "Bolívar", "FC The Strongest": "The Strongest", "C Jorge Wilstermann": "Jorge Wilstermann",
           "C Oriente Petrolero": "Oriente Petrolero", "Blooming FC": "Blooming", "C Aurora": "Aurora",
           "C Guabirá": "Guabirá", "C Independiente Petrolero": "Independiente Petrolero", "C San José": "San José"},
    "VE": {"Caracas FC": "Caracas", "Deportivo Táchira FC": "Deportivo Táchira", "Carabobo FC": "Carabobo",
           "Monagas SC": "Monagas", "Estudiantes de Mérida FC": "Estudiantes de Mérida", "Portuguesa FC": "Portuguesa",
           "UC de Venezuela": "Universidad Central"},
}
HEAD = PatternFill("solid", fgColor="1F4E78")
MARKED = PatternFill("solid", fgColor="E2EFDA")

LEIA = """Clubes da CONMEBOL de 2026 (Argentina, Uruguai, Colômbia, Paraguai, Equador, Chile, Peru, Bolívia, Venezuela)

1. ABA CLUBES: todos os clubes desses países que existem no FM06 (o jogo não cria clube novo).
   - usar = sim: o clube entra com o elenco de 2026. Os que acho que estão na primeira divisão de 2026 já vêm com
     "sim" e o nome de 2026 (fundo verde): confira, tire o "sim" de quem não está e ponha nos que faltam.
   - clube_2026: o nome que o mod usa, com a sigla do país no fim ("River Plate (ARG)"; se escrever sem, ela entra
     sozinha). É o mesmo nome da coluna clube_2026 das abas Jogadores e Carreiras; lá pode ir sem a sigla quando
     só um clube marcado tem o nome ("Boca Juniors"), mas River Plate, Nacional e Universidad Católica existem em
     dois países e precisam dela.
   - tecnico, tecnico_nacionalidade (ARG, URU, BRA, ESP...), tecnico_nascimento (AAAA-MM-DD): o técnico de 2026.
     Ele ocupa o técnico do clube no banco de 2005 (o resto da comissão fica como está).
   - jogadores_no_fm06: quantos jogadores o clube tem no banco de 2005 (só para referência; se o elenco de 2026 tiver
     mais, o mod traz jogadores de clubes que não usamos).
   - Clube de 2026 que não existe no FM06 (Tigre, Sarmiento, Independiente del Valle, Always Ready...) fica de fora.
   - Não mexa em pais, clube_fm06, nome_curto_fm06 e uid_jogo.

2. ABA JOGADORES: uma linha por jogador de 2026 de cada clube marcado (até 45 por clube).
   Obrigatórios:
   - clube_2026: igual ao da aba Clubes ("River Plate (ARG)");
   - nome: nome_comum (como aparece no jogo: "Franco Armani") e/ou primeiro_nome + sobrenome;
   - nascimento: AAAA-MM-DD (ou data do Excel);
   - posicoes: uma ou mais, separadas por espaço, na ordem de preferência: GK CB RB LB DM CM AM RW LW ST.
   Opcionais (vazio = estimado):
   - nacionalidade: em português (Argentina, Uruguai, Colômbia, Paraguai, Equador, Chile, Peru, Bolívia,
     Venezuela, Brasil, Espanha, Itália...). Vazio = o país do clube;
   - pe: D, E ou A (ambos); altura em cm;
   - valor_mercado_eur: valor de mercado em euros (o Transfermarkt serve). É o que mais pesa no CA;
   - ca e pa (1 a 200): se quiser dar a força à mão; vazio = calculado pelo valor, pela idade e pelo país
     (Argentina como a Série A; Uruguai, Colômbia, Paraguai, Equador e Chile como a B; Peru, Bolívia e
     Venezuela como a C);
   - contrato_fim (AAAA-MM-DD; vazio = 2026-12-31), numero (camisa), jogos_selecao e gols_selecao.

3. ABA CARREIRAS (opcional): os clubes por onde o jogador passou, para a aba Histórico.
   - clube_2026, nome_comum e nascimento iguais aos da aba Jogadores;
   - temporada: 2019, 2019–2021 ou 2024– (em aberto); clube: o nome do clube ("Boca Juniors", "Palmeiras",
     "Real Madrid"); emprestimo: sim ou vazio; valor_eur: o valor pago para chegar (opcional).
   Sem carreira, o jogador fica só com a temporada atual.

Depois de preencher: python scripts\\estrangeiros_planilha.py (confere e monta data\\estrangeiros) e
python scripts\\aplicar_tudo.py (completo, sem --rapido).

Exemplo de linha da aba Jogadores:
River Plate (ARG) | 1 | Franco Armani | Franco | Armani | 1986-10-16 | Argentina | GK | D | 189 | 1000000 | | | 2026-12-31 | 21 | 0
"""


def _contracts() -> Counter:
    if not PEOPLE_BACKUP.exists():
        return Counter()
    return Counter(c.club for c in people_contracts.contracts(PEOPLE_BACKUP.read_bytes()).values())


def _header(ws, columns: list[str], widths: dict[str, int]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        ws.column_dimensions[cell.column_letter].width = widths.get(cell.value, 14)
    ws.freeze_panes = "A2"


def write_model(path: Path) -> None:
    server = server_bytes()
    brazil = set(people_contracts.brazil_club_indexes(server))
    count = _contracts()
    book = openpyxl.Workbook()
    leia = book.active
    leia.title = "LEIA"
    for line in LEIA.splitlines():
        leia.append([line])
    leia.column_dimensions["A"].width = 130
    clubs = book.create_sheet("Clubes")
    _header(clubs, CLUB_COLUMNS, {"clube_fm06": 36, "nome_curto_fm06": 22, "clube_2026": 26, "tecnico": 24,
                                  "tecnico_nacionalidade": 12, "tecnico_nascimento": 14, "fonte": 30, "obs": 30})
    marked = 0
    for alpha2, (_, code, _, _) in ex.NATIONS.items():
        first = FIRST_2026.get(alpha2, {})
        found = {c["uid"]: c for c in ex.fm06_clubs(server, alpha2, brazil)}.values()
        for club in sorted(found, key=lambda c: (c["long"] not in first, -count[c["index"]], c["long"])):
            name = ex.display_name(first[club["long"]], alpha2) if club["long"] in first else ""
            clubs.append([code, club["long"], club["short"], club["uid"], count[club["index"]],
                          "sim" if name else "", name])
            if name:
                marked += 1
                for cell in clubs[clubs.max_row]:
                    cell.fill = MARKED
    yes = DataValidation(type="list", formula1='"sim,não"', allow_blank=True)
    clubs.add_data_validation(yes)
    yes.add(f"F2:F{clubs.max_row}")
    players = book.create_sheet("Jogadores")
    _header(players, PLAYER_COLUMNS, {"clube_2026": 24, "nome_comum": 24, "primeiro_nome": 16, "sobrenome": 18,
                                      "posicoes": 14, "fonte": 30, "obs": 30})
    foot = DataValidation(type="list", formula1='"D,E,A"', allow_blank=True)
    players.add_data_validation(foot)
    foot.add("I2:I3000")
    careers = book.create_sheet("Carreiras")
    _header(careers, CAREER_COLUMNS, {"clube_2026": 24, "nome_comum": 24, "clube": 28, "fonte": 30})
    for sheet in (leia, clubs, players, careers):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    print(f"[OK] {path.name}: {clubs.max_row - 1} clubes do FM06, {marked} marcados como da primeira "
          "divisão de 2026")


def _rows(book, name: str) -> list[dict]:
    rows = book[name].iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    out = []
    for values in rows:
        row = {h: (v.strip() if isinstance(v, str) else v) for h, v in zip(header, values) if h}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)
    return out


def _date(value) -> str:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()[:10]
    try:
        return dt.date.fromisoformat(text).isoformat()
    except ValueError:
        return ""


def _int(value) -> int | None:
    text = str(value if value is not None else "").replace(".", "").replace(",", "").strip()
    return int(float(text)) if text.replace("-", "").isdigit() else None


def build(model: Path) -> list[str]:
    """Monta data/estrangeiros; devolve os erros (linhas que ficaram de fora)."""
    book = openpyxl.load_workbook(model, read_only=True, data_only=True)
    errors: list[str] = []
    server = server_bytes()
    brazil = set(people_contracts.brazil_club_indexes(server))
    repo = {}
    for alpha2 in ex.NATIONS:
        for club in ex.fm06_clubs(server, alpha2, brazil):
            repo.setdefault(club["uid"], club["repo"])

    clubs, coaches, names, bare = [], [], {}, {}
    for line, row in enumerate(_rows(book, "Clubes"), start=2):
        if str(row.get("usar") or "").lower() not in ("sim", "s", "x"):
            continue
        alpha2, uid, name = BY_CODE.get(str(row.get("pais") or "").upper()), _int(row.get("uid_jogo")), row.get("clube_2026")
        if not alpha2 or uid is None or not name:
            errors.append(f"Clubes linha {line}: usar=sim sem pais, uid_jogo ou clube_2026")
            continue
        full = name if re.search(r"\([A-Z]{3}\)$", name) else ex.display_name(name, alpha2)
        if full in names:
            errors.append(f"Clubes linha {line}: clube_2026 '{full}' repetido")
            continue
        names[full] = (full, alpha2)
        bare.setdefault(re.sub(r"\s*\([A-Z]{3}\)$", "", full), []).append(full)
        clubs.append({"clube_2026": full, "pais": alpha2, "uid_jogo": uid, "uid_repo": repo.get(uid, ""),
                      "fm06_nome": row.get("nome_curto_fm06") or "", "tipo": "estrangeiro", "onde": "planilha",
                      "sofascore_id": "", "sofascore_nome": name})
        if row.get("tecnico"):
            born = _date(row.get("tecnico_nascimento"))
            if row.get("tecnico_nascimento") and not born:
                errors.append(f"Clubes linha {line}: tecnico_nascimento '{row['tecnico_nascimento']}' não é AAAA-MM-DD")
            coaches.append({"clube_2026": full, "pais": alpha2, "cargo": "tecnico", "nome": row["tecnico"],
                            "nacionalidade": str(row.get("tecnico_nacionalidade") or "").upper() or ex.NATIONS[alpha2][1],
                            "nascimento": born, "fonte": row.get("fonte") or "planilha"})

    for key, full in bare.items():      # "Boca Juniors" sem o país, quando só um clube tem o nome
        if len(full) == 1:
            names.setdefault(key, names[full[0]])
    players, seen = [], Counter()
    for line, row in enumerate(_rows(book, "Jogadores"), start=2):
        club = names.get(row.get("clube_2026"))
        common = row.get("nome_comum") or ""
        first, second = row.get("primeiro_nome") or "", row.get("sobrenome") or ""
        if not common and not (first or second):
            errors.append(f"Jogadores linha {line}: sem nome")
            continue
        if not first and not second:
            first, _, second = common.partition(" ")
        label = f"Jogadores linha {line} ({common or first + ' ' + second})"
        if club is None:
            errors.append(f"{label}: clube_2026 '{row.get('clube_2026')}' não está marcado (usar=sim) na aba Clubes")
            continue
        born = _date(row.get("nascimento"))
        roles = str(row.get("posicoes") or "").upper().replace(",", " ").split()
        if not born or not roles or any(r not in ROLE_POSITIONS for r in roles):
            errors.append(f"{label}: nascimento (AAAA-MM-DD) e posicoes ({' '.join(ROLE_POSITIONS)}) obrigatórios")
            continue
        seen[club[0]] += 1
        if seen[club[0]] > 45:
            errors.append(f"{label}: mais de 45 jogadores no clube; fica de fora")
            continue
        alpha2 = club[1]
        serie = ex.NATIONS[alpha2][3]
        nation = row.get("nacionalidade") or ex.NATIONS[alpha2][2]
        if nation not in people_contracts.NATION_CODES:
            errors.append(f"{label}: nacionalidade '{nation}' sem código no jogo (fica a do jogador de 2005)")
        height = _int(row.get("altura")) or mj.HEIGHT.get(roles[0], 178)
        value = _int(row.get("valor_mercado_eur"))
        age = mj._age(dt.date.fromisoformat(born))
        caps = _int(row.get("jogos_selecao"))
        ca = _int(row.get("ca")) or mj.estimate_ca(value, serie, age, False, caps)
        pa = max(ca, _int(row.get("pa")) or mj.estimate_pa(ca, value, serie, age, hash((common, born)) % 1000))
        estimated = [n for n, v in (("altura", row.get("altura")), ("pé", row.get("pe")), ("CA", row.get("ca")),
                                    ("PA", row.get("pa"))) if not v] + ["peso", "salário"]
        players.append({
            "clube": club[0], "numero": _int(row.get("numero")), "nome_comum": common, "primeiro_nome": first,
            "sobrenome": second, "nascimento": born, "nacionalidade": nation, "segunda_nacionalidade": None,
            "posicoes": " ".join(roles), "pe": str(row.get("pe") or "D").upper()[:1], "altura": height,
            "peso": round((24.0 if roles[0] == "GK" else 23.2) * (height / 100) ** 2), "valor_mercado_eur": value,
            "ca": min(200, ca), "pa": min(200, pa), "contrato_inicio": None,
            "contrato_fim": _date(row.get("contrato_fim")) or "2026-12-31",
            "salario_mensal_brl": mj.estimate_salary(value, serie, False), "situacao_elenco": None,
            "emprestado_de": None, "emprestimo_fim": None, "jogos_selecao": caps,
            "gols_selecao": _int(row.get("gols_selecao")), "destaques": "", "fonte_url": row.get("fonte") or "planilha",
            "obs": "estimado: " + ", ".join(estimated)})
    mj._squad_roles(players)

    careers = []
    by_player = {(p["clube"], p["nome_comum"] or f"{p['primeiro_nome']} {p['sobrenome']}".strip(), p["nascimento"])
                 for p in players}
    for line, row in enumerate(_rows(book, "Carreiras"), start=2):
        club = names.get(row.get("clube_2026"))
        key = (club[0] if club else "", row.get("nome_comum") or "", _date(row.get("nascimento")))
        if key not in by_player:
            errors.append(f"Carreiras linha {line}: {row.get('nome_comum')} ({row.get('clube_2026')}, "
                          f"{row.get('nascimento')}) não está na aba Jogadores com esse clube, nome e nascimento")
            continue
        if not row.get("temporada") or not row.get("clube"):
            errors.append(f"Carreiras linha {line}: sem temporada ou clube")
            continue
        careers.append({"clube_2026": key[0], "nome_comum": key[1], "nascimento": key[2],
                        "temporada": str(row["temporada"]).replace("-", "–"), "clube": row["clube"],
                        "emprestimo": "sim" if str(row.get("emprestimo") or "").lower() in ("sim", "s", "x") else "não",
                        "jogos": "", "gols": "", "valor_eur": _int(row.get("valor_eur")) or "",
                        "fonte": row.get("fonte") or "planilha", "atualizado": dt.date.today().isoformat()})

    ex.DATA.mkdir(parents=True, exist_ok=True)
    with ex.IDENTITIES.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=ex.FIELDS + ["manual"])
        writer.writeheader()
        writer.writerows({**c, "manual": "sim"} for c in clubs)
    with ex.COACHES.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=COACH_FIELDS)
        writer.writeheader()
        writer.writerows(coaches)
    with ex.CAREERS.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=CAREER_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(careers)
    sheet_book = openpyxl.Workbook()
    sheet = sheet_book.active
    sheet.title = "Jogadores"
    header = [k for k in players[0]] if players else ["clube"]
    sheet.append(header)
    for p in players:
        sheet.append([p.get(h) for h in header])
    sheet_book.save(ex.PLAYERS)

    per_club = Counter(p["clube"] for p in players)
    empty = [c["clube_2026"] for c in clubs if not per_club[c["clube_2026"]]]
    print(f"[OK] {len(clubs)} clubes, {len(players)} jogadores, {len(coaches)} técnicos e {len(careers)} passagens "
          f"em {ex.DATA}")
    if empty:
        print(f"[AVISO] {len(empty)} clubes marcados sem jogadores (ficam com o elenco de 2005): "
              + ", ".join(empty[:12]) + (" ..." if len(empty) > 12 else ""))
    return errors


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--modelo", action="store_true", help="cria a planilha (não sobrescreve uma preenchida)")
    parser.add_argument("--planilha", type=Path, default=MODEL)
    args = parser.parse_args()
    if args.modelo:
        if args.planilha.exists():
            sys.exit(f"[ERRO] {args.planilha} já existe; apague ou renomeie antes de criar outra.")
        write_model(args.planilha)
        return
    errors = build(args.planilha)
    for error in errors:
        print("[CONFIRA]", error)
    if errors:
        print(f"[CONFIRA] {len(errors)} linhas com problema (as outras entraram); corrija e rode de novo.")
    print("Depois: python scripts\\aplicar_tudo.py (completo, sem --rapido).")


if __name__ == "__main__":
    main()
