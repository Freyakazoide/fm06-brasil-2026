"""Monta data/brasil/jogadores/jogadores_2026.xlsx com os dados do Sofascore.

Uso: python scripts/montar_jogadores.py [--cache workspace/sofascore]
(depois de scripts/coletar_sofascore.py)

Dados reais: nome, número, posições, nascimento, altura, pé, nacionalidade, fim
de contrato, valor de mercado, início do contrato e empréstimo (transferências),
jogos e gols pela seleção. Estimados (marcados na coluna obs): peso (pela
altura), salário (~10% do valor por ano, piso por série), o que faltar no
Sofascore e CA/PA:

- CA pelo valor de mercado: 104 + 22 × log10(valor / 100 mil), menos um
  desconto por série (o mesmo valor pesa menos numa série menor), mais
  ajuste pela idade (jovem vale mais do que joga, veterano menos; queda
  depois dos 35) e bônus pelos jogos na seleção principal. Sem valor: um CA
  típico de quem não tem valor na série.
- PA = CA + crescimento pela idade, maior para jovem de valor alto.
- Destaques: a visão geral do Sofascore (ataque, técnica, tática, defesa,
  criatividade; goleiro: defesas, antecipação, distribuição, jogo aéreo)
  desloca os atributos ligados a ela em relação à média do próprio jogador.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from fm06patch.squad_builder import Target, estimate_attributes  # noqa: E402

OUT = ROOT / "data" / "brasil" / "jogadores" / "jogadores_2026.xlsx"
TEMPLATE = ROOT / "data" / "brasil" / "jogadores" / "jogadores_2026_modelo.xlsx"
START = date(2026, 1, 1)

POSITIONS = {"GK": "GK", "DC": "CB", "DR": "RB", "DL": "LB", "DM": "DM", "MC": "CM", "AM": "AM",
             "ML": "LW", "MR": "RW", "LW": "LW", "RW": "RW", "ST": "ST"}
BY_LETTER = {"G": "GK", "D": "CB", "M": "CM", "F": "ST"}
FOOT = {"Right": "D", "Left": "E", "Both": "A"}
HEIGHT = {"GK": 189, "CB": 186, "ST": 182}
# desconto do CA pelo valor em cada série; CA de quem não tem valor
SERIE_DISCOUNT = {"A": 0, "B": 3, "C": 8, "D": 14}
NO_VALUE_CA = {"A": 118, "B": 100, "C": 88, "D": 76}
YOUTH_CA = {"A": 82, "B": 74, "C": 66, "D": 58}
SALARY_FLOOR = {"A": 30_000, "B": 15_000, "C": 6_000, "D": 2_500}   # R$ por mês
EUR_BRL = 5.8
MAX_AGE = 47                                            # mais velho que isso: erro de cadastro ou comissão (Fábio tem 45)
YOUTH_LIMIT = 30                                        # sub-20 por clube (os de maior PA)
GROWTH = {15: 40, 16: 38, 17: 35, 18: 30, 19: 26, 20: 22, 21: 18, 22: 14, 23: 10, 24: 7, 25: 4, 26: 2}
COUNTRIES = {
    "Brazil": "Brasil", "Argentina": "Argentina", "Colombia": "Colômbia", "Uruguay": "Uruguai",
    "Paraguay": "Paraguai", "Portugal": "Portugal", "Ecuador": "Equador", "Venezuela": "Venezuela",
    "Italy": "Itália", "Chile": "Chile", "Peru": "Peru", "Nigeria": "Nigéria", "Spain": "Espanha",
    "Cameroon": "Camarões", "Panama": "Panamá", "Angola": "Angola", "Ghana": "Gana", "Bolivia": "Bolívia",
    "Morocco": "Marrocos", "Cape Verde": "Cabo Verde", "Croatia": "Croácia", "USA": "Estados Unidos",
    "Guinea": "Guiné", "South Korea": "Coreia do Sul", "Netherlands": "Holanda", "England": "Inglaterra",
    "Sweden": "Suécia", "Denmark": "Dinamarca", "DR Congo": "RD Congo", "Senegal": "Senegal",
    "Guinea-Bissau": "Guiné-Bissau", "Bulgaria": "Bulgária", "Saudi Arabia": "Arábia Saudita",
    "Mexico": "México", "Anguilla": "Anguila", "Azerbaijan": "Azerbaijão", "Côte d'Ivoire": "Costa do Marfim",
    "Haiti": "Haiti", "Fiji": "Fiji", "France": "França", "Germany": "Alemanha", "Japan": "Japão",
    "Costa Rica": "Costa Rica", "Honduras": "Honduras", "Mozambique": "Moçambique", "Belgium": "Bélgica",
}
OVERVIEW = {
    "attacking": ("Finalização", "Sem Bola", "Compostura"),
    "technical": ("Técnica", "Primeiro Toque", "Finta"),
    "tactical": ("Posicionamento", "Decisões", "Trabalho Equipa"),
    "defending": ("Marcação", "Desarme", "Antecipação"),
    "creativity": ("Criatividade", "Passe"),
    "saves": ("Reflexes", "One on one", "Handling"),
    "anticipation": ("Antecipação", "Rushing out"),
    "ballDistribution": ("Kicking", "Passe"),
    "aerial": ("Aer ability", "Cmd of area", "Punch"),
}


class Cache:
    def __init__(self, path: Path):
        self.path = path

    def get(self, request: str):
        key = self.path / (hashlib.md5(request.encode()).hexdigest() + ".json")
        return json.loads(key.read_text()) if key.exists() else None


def _date(timestamp) -> date | None:
    return datetime.fromtimestamp(timestamp, timezone.utc).date() if timestamp else None


def _age(born: date) -> float:
    return (START - born).days / 365.25


def estimate_ca(value, serie: str, age: float, youth: bool, caps: int | None = None) -> int:
    if value:
        ca = 104 + 22 * math.log10(value / 100_000) - SERIE_DISCOUNT[serie]
        ca -= max(0.0, 23 - age) * 2.5
    elif youth:
        ca = YOUTH_CA[serie] + (min(age, 20) - 17) * 4
    else:
        ca = NO_VALUE_CA[serie] - max(0.0, 23 - age) * 4
    if age > 31:                                        # veterano vale menos do que joga; depois dos 35 cai
        ca += 2.5 * min(age - 31, 4) - 3 * max(0.0, age - 35)
    ca += min(8, (caps or 0) / 8)                       # seleção principal
    return int(max(40, min(172, round(ca))))


def estimate_pa(ca: int, value, serie: str, age: float, seed: int = 0) -> int:
    growth = GROWTH.get(int(age), 0 if age > 26 else 40)
    if age < 23:                                        # nem toda promessa vinga: ±8 fixo por jogador
        growth += seed % 17 - 8
    if value and age < 24:
        relative = math.log10(value / 100_000) - SERIE_DISCOUNT[serie] / 22
        growth += max(0, min(12, round(6 * relative)))
    return int(max(ca, min(190, ca + growth)))


def estimate_salary(value, serie: str, youth: bool) -> int:
    """R$ por mês: ~10% do valor de mercado por ano, com piso por série."""
    floor = 1_500 if youth else SALARY_FLOOR[serie]
    salary = (value or 0) * 0.10 * EUR_BRL / 12
    return int(round(max(floor, salary), -2))


def _names(player: dict) -> tuple[str, str, str]:
    """(nome comum, primeiro nome, sobrenome)."""
    full = (player.get("name") or "").replace('"', "").strip()      # aspas quebram o .lnc
    short = (player.get("shortName") or "").replace('"', "").strip()
    parts = full.split() or [short]
    if len(parts) == 1:
        return parts[0], parts[0], ""
    common = short if short and "." not in short and short != full else ""
    return common, parts[0], " ".join(parts[1:])


def _positions(player: dict) -> tuple[list[str], bool]:
    detailed = [POSITIONS[p] for p in (player.get("positionsDetailed") or []) if p in POSITIONS]
    roles = list(dict.fromkeys(detailed))[:3]
    if roles:
        return roles, False
    return [BY_LETTER.get(player.get("position") or "M", "CM")], True


def _contract(cache: Cache, player: dict, team_ids: set[int]):
    """(início, emprestado de) pela última transferência para o clube."""
    history = (cache.get(f"/player/{player['id']}/transfer-history") or {}).get("transferHistory", [])
    moves = sorted((t for t in history if (t.get("transferTo") or {}).get("id") in team_ids),
                   key=lambda t: t.get("transferDateTimestamp") or 0)
    if not moves:
        return None, None
    last = moves[-1]
    loan = last.get("type") == 1
    return _date(last.get("transferDateTimestamp")), (last.get("transferFrom") or {}).get("name") if loan else None


def _highlights(cache: Cache, player: dict, target: Target) -> str:
    data = cache.get(f"/player/{player['id']}/attribute-overviews") or {}
    current = next((o for o in data.get("playerAttributeOverviews", []) if o.get("yearShift") == 0), None)
    if not current:
        return ""
    scores = {k: v for k, v in current.items() if k in OVERVIEW and isinstance(v, (int, float))}
    if target.is_goalkeeper:
        scores.pop("tactical", None)
    if not scores:
        return ""
    mean = sum(scores.values()) / len(scores)
    base = estimate_attributes(target)
    out = {}
    for key, score in scores.items():
        for name in OVERVIEW[key]:
            if name in base:
                out[name] = max(1, min(20, round(base[name] + (score - mean) / 5)))
    return "; ".join(f"{k}={v}" for k, v in out.items())


def _caps(cache: Cache, player: dict) -> tuple[int | None, int | None]:
    stats = (cache.get(f"/player/{player['id']}/national-team-statistics") or {}).get("statistics", [])
    senior = [s for s in stats if "U2" not in s["team"]["name"] and "U1" not in s["team"]["name"]]
    if not senior:
        return None, None
    return sum(s.get("appearances", 0) for s in senior), sum(s.get("goals", 0) for s in senior)


def player_row(cache: Cache, item: dict, club: dict, youth: bool) -> dict:
    player = item["player"]
    estimated = []
    common, first, last = _names(player)
    roles, guessed = _positions(player)
    if guessed:
        estimated.append("posição")
    born = _date(player.get("dateOfBirthTimestamp"))
    if born is None:
        born = date(2007 if youth else 2000, 1, 1)
        estimated.append("nascimento")
    height = player.get("height")
    if not height:
        height = HEIGHT.get(roles[0], 178)
        estimated.append("altura")
    foot = FOOT.get(player.get("preferredFoot") or "")
    if not foot:
        foot = "D"
        estimated.append("pé")
    weight = round((24.0 if roles[0] == "GK" else 23.2) * (height / 100) ** 2)
    estimated += ["peso", "salário"]
    value = player.get("proposedMarketValue")
    age = _age(born)
    caps, goals = _caps(cache, player)
    ca = estimate_ca(value, club["serie"], age, youth, caps)
    pa = estimate_pa(ca, value, club["serie"], age, player["id"])
    team_ids = {club["sofascore_id"], club.get("sub20_id")}
    started, loaned_from = _contract(cache, player, team_ids)
    until = _date(player.get("contractUntilTimestamp"))
    if until is None:
        until = date(2026, 12, 31)
        estimated.append("fim do contrato")
    country = (player.get("country") or {}).get("name")
    target = Target(number="", common=common, first=first, second=last, born=born.isoformat(),
                    nationality="", roles=roles, foot=foot, height=height, weight=weight, ca=ca, pa=pa)
    return {
        "clube": club["clube"], "numero": player.get("shirtNumber") or player.get("jerseyNumber"),
        "nome_comum": common, "primeiro_nome": first, "sobrenome": last, "nascimento": born.isoformat(),
        "nacionalidade": COUNTRIES.get(country, country or "Brasil"), "segunda_nacionalidade": None,
        "posicoes": " ".join(roles), "pe": foot, "altura": height, "peso": weight,
        "valor_mercado_eur": value, "ca": ca, "pa": pa,
        "contrato_inicio": started.isoformat() if started else None, "contrato_fim": until.isoformat(),
        "salario_mensal_brl": estimate_salary(value, club["serie"], youth), "situacao_elenco": "Base (sub-20)" if youth else None,
        "emprestado_de": loaned_from, "emprestimo_fim": until.isoformat() if loaned_from else None,
        "jogos_selecao": caps, "gols_selecao": goals,
        "destaques": _highlights(cache, player, target),
        "fonte_url": f"https://www.sofascore.com/player/{player.get('slug', '')}/{player['id']}",
        "obs": ("estimado: " + ", ".join(estimated)) if estimated else "",
        "_serie": club["serie"], "_id": player["id"],
    }


def _squad_roles(rows: list[dict]) -> None:
    senior = sorted((r for r in rows if not r["situacao_elenco"]), key=lambda r: -r["ca"])
    for rank, row in enumerate(senior):
        young = _age(date.fromisoformat(row["nascimento"])) < 21
        row["situacao_elenco"] = ("Titular" if rank < 11 else "Rotação" if rank < 20
                                  else "Jovem promessa" if young else "Reserva")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache", type=Path, default=ROOT / "workspace" / "sofascore")
    args = parser.parse_args()
    cache = Cache(args.cache)
    clubs = json.loads((args.cache / "times.json").read_text(encoding="utf-8"))
    rows: list[dict] = []
    listed = {x["player"]["id"] for c in clubs for tid in (c["sofascore_id"], c.get("sub20_id")) if tid
              for x in (cache.get(f"/team/{tid}/players") or {}).get("players", [])}
    for club in clubs:
        seen: set[int] = set()
        club_rows = []
        for tid, youth in ((club["sofascore_id"], False), (club.get("sub20_id"), True)):
            if not tid:
                continue
            items = list((cache.get(f"/team/{tid}/players") or {}).get("players", []))
            if not youth:                               # elenco que o Sofascore não lista: escalações de 2026
                items += [p for p in (cache.get(f"/player/{pid}") for pid in club.get("escalados", [])
                                      if pid not in listed) if p]
            for item in items:
                born = _date(item["player"].get("dateOfBirthTimestamp"))
                if item["player"]["id"] in seen or (born and _age(born) > MAX_AGE):
                    continue
                seen.add(item["player"]["id"])
                club_rows.append(player_row(cache, item, club, youth))
        youth_rows = sorted((r for r in club_rows if r["situacao_elenco"]), key=lambda r: -r["pa"])
        club_rows = [r for r in club_rows if not r["situacao_elenco"]] + youth_rows[:YOUTH_LIMIT]
        _squad_roles(club_rows)
        rows += club_rows
        if len(club_rows) < 18:
            print(f"[AVISO] {club['clube']}: só {len(club_rows)} jogadores no Sofascore")

    book = openpyxl.load_workbook(TEMPLATE)
    sheet = book["Jogadores"]
    header = [c.value for c in sheet[1]]
    sheet.delete_rows(2, sheet.max_row)
    for row in rows:
        sheet.append([row.get(h) for h in header])
    values = book["Valores"]
    values["D6"] = "Base (sub-20)"
    book.save(OUT)
    series = {s: sum(1 for r in rows if r["_serie"] == s) for s in "ABCD"}
    print(f"[OK] {len(rows)} jogadores de {len(clubs)} clubes em {OUT} ({series})")


if __name__ == "__main__":
    main()
