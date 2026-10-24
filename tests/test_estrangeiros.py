from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from fm06patch import estrangeiros as ex  # noqa: E402
from fm06patch import historico_jogadores as hj  # noqa: E402

RIVER = {"uid": 94, "repo": 93, "index": 33, "long": "CA River Plate", "short": "River"}
CLUBS = [RIVER,
         {"uid": 82, "repo": 81, "index": 22, "long": "CA Boca Juniors", "short": "Boca"},
         {"uid": 85, "repo": 83, "index": 24, "long": "Estudiantes de La Plata", "short": "Estudiantes (LP)"},
         {"uid": 86, "repo": 84, "index": 25, "long": "Estudiantes de Buenos Aires", "short": "Estudiantes (BA)"},
         {"uid": 87, "repo": 86, "index": 26, "long": "Central Córdoba de Rosario", "short": "Ctral. Córdoba"},
         {"uid": 88, "repo": 87, "index": 27, "long": "Huracán", "short": "Huracán"},
         {"uid": 89, "repo": 88, "index": 28, "long": "Huracán de Tres Arroyos", "short": "Huracán (TA)"},
         {"uid": 90, "repo": 89, "index": 29, "long": "CCD Tolima", "short": "Tolima"},
         {"uid": 91, "repo": 90, "index": 30, "long": "Lotería del Tolima Cooperamos", "short": "Cooperamos Tolima"},
         {"uid": 92, "repo": 91, "index": 31, "long": "C Deportivo Recoleta", "short": "Deportivo Recoleta"}]


class MatchClubTest(unittest.TestCase):
    def test_full_name_without_prefixes(self):
        self.assertEqual(ex.match_club("River Plate", CLUBS)["uid"], 94)
        self.assertEqual(ex.match_club("Boca Juniors", CLUBS)["uid"], 82)

    def test_contained_name_only_when_unique(self):
        self.assertEqual(ex.match_club("Estudiantes de La Plata", CLUBS)["uid"], 85)
        self.assertIsNone(ex.match_club("Estudiantes", CLUBS))         # dois clubes com o nome
        self.assertIsNone(ex.match_club("Deportivo Riestra", CLUBS))   # não existe no FM06

    def test_alias(self):
        self.assertEqual(ex.match_club("CARP", CLUBS, {"CARP": 94})["uid"], 94)
        self.assertEqual(ex.match("Gimnasia", CLUBS, {"Gimnasia": "Estudiantes de Buenos Aires"}), (CLUBS[3], "apelido"))

    def test_another_city_is_another_club(self):
        self.assertIsNone(ex.match_club("Central Córdoba", CLUBS))     # o de Santiago del Estero
        self.assertEqual(ex.match("Huracán", CLUBS), (CLUBS[5], "nome"))    # não o de Tres Arroyos

    def test_prefix_or_legal_form_more(self):
        self.assertEqual(ex.match("Deportes Tolima", CLUBS), (CLUBS[7], "aproximado"))
        self.assertEqual(ex.match("Recoleta", CLUBS), (CLUBS[9], "aproximado"))

    def test_display_name_has_the_country(self):
        self.assertEqual(ex.display_name("Nacional", "UY"), "Nacional (URU)")


class IdentitiesTest(unittest.TestCase):
    def setUp(self) -> None:
        self.dir = Path(tempfile.mkdtemp())
        self.br = self.dir / "identidade_fm06.csv"
        self.br.write_text("clube_2026,uid_jogo,uid_repo,tipo\nPalmeiras,329,328,clube\n", encoding="utf-8-sig")
        self.abroad = self.dir / "identidade_estrangeiros.csv"
        with self.abroad.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=ex.FIELDS)
            writer.writeheader()
            writer.writerow({"clube_2026": "River Plate (ARG)", "pais": "AR", "uid_jogo": 94, "uid_repo": 93,
                             "tipo": "estrangeiro"})
            writer.writerow({"clube_2026": "Riestra (ARG)", "pais": "AR", "uid_jogo": "", "tipo": "estrangeiro"})

    def test_rows_without_fm06_club_stay_out(self):
        with mock.patch.object(ex, "IDENTITIES", self.abroad):
            self.assertEqual([r["clube_2026"] for r in ex.identity_rows()], ["River Plate (ARG)"])
            self.assertEqual(ex.nation_of(94), 0xBB)
            self.assertEqual(ex.nation_of(329), 0xBD)

    def test_counting_list_has_brazil_and_abroad(self):
        out = self.dir / "todas.csv"
        with mock.patch.object(ex, "IDENTITIES", self.abroad), mock.patch.object(ex, "BR_IDENTITIES", self.br):
            ex.write_all_identities(out)
        rows = list(csv.DictReader(out.open(encoding="utf-8-sig")))
        self.assertEqual([(r["clube_2026"], r["uid_jogo"]) for r in rows], [("Palmeiras", "329"), ("River Plate (ARG)", "94")])

    def test_no_foreign_file(self):
        with mock.patch.object(ex, "IDENTITIES", self.dir / "nada.csv"):
            self.assertEqual(ex.identity_rows(), [])


class CareerClubsTest(unittest.TestCase):
    def test_foreign_2026_name_finds_the_club(self):
        names = {94: ("CA River Plate", "River"), 329: ("SE Palmeiras", "Palmeiras")}
        clubs = hj.Clubs(names, hj.Counter(), {"Palmeiras": 329, "River Plate (ARG)": 94}, strict=True)
        self.assertEqual(clubs.find("River Plate (ARG)"), (94, "brasil"))


BASELINE = ROOT / "workspace" / "baseline" / "server_db.original.dat"


@unittest.skipUnless(BASELINE.exists(), "precisa do server_db original em workspace/baseline")
class RealDatabaseTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.server = BASELINE.read_bytes()

    def test_clubs_of_each_country(self):
        argentina = {c["uid"]: c for c in ex.fm06_clubs(self.server, "AR")}
        self.assertEqual((argentina[94]["index"], argentina[82]["short"]), (33, "Boca"))
        for alpha2 in ex.NATIONS:                   # sem os pares falsos (belgas, brasileiros, UID 0)
            for club in ex.fm06_clubs(self.server, alpha2):
                self.assertTrue(0 < club["repo"] < club["uid"], club)
        self.assertIn("Colo Colo", {c["short"] for c in ex.fm06_clubs(self.server, "CL")})
        bolivia = {c["short"]: c["index"] for c in ex.fm06_clubs(self.server, "BO")}
        self.assertEqual((bolivia["Bolívar"], bolivia["The Strongest"]), (188, 189))   # índice = código de nação

    def test_indexes_and_team_values(self):
        rows = [{"clube_2026": "River Plate (ARG)", "pais": "AR", "uid_jogo": "94"},
                {"clube_2026": "Peñarol (URU)", "pais": "UY", "uid_jogo": "1922"}]
        self.assertEqual(ex.club_indexes(self.server, rows), {94: 33, 1922: 1467})
        values = ex.team_values(self.server)
        self.assertEqual((values[94], values[1922], values[329]), (84, 1665, 297))

    def test_montar_estrangeiros_matches_by_name(self):
        import montar_estrangeiros
        teams = [{"pais": "AR", "clube": "River Plate (ARG)", "liga": "Liga", "sofascore_id": 1, "sofascore_nome": "River Plate"},
                 {"pais": "AR", "clube": "Riestra (ARG)", "liga": "Liga", "sofascore_id": 2, "sofascore_nome": "Riestra"},
                 {"pais": "UY", "clube": "Peñarol (URU)", "liga": "Liga", "sofascore_id": 3, "sofascore_nome": "Peñarol"}]
        kept = {2: {"uid_jogo": "82", "uid_repo": "81", "fm06_nome": "Boca"}}     # posto à mão
        rows, notes = montar_estrangeiros.identities(teams, self.server, kept)
        self.assertEqual([r["uid_jogo"] for r in rows], [94, 82, 1922])
        self.assertEqual(notes, [])


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(BASELINE.exists(), "precisa do server_db original em workspace/baseline")
class SpreadsheetTest(unittest.TestCase):
    def test_model_to_mod_files(self):
        import openpyxl
        import estrangeiros_planilha as sheet
        tmp = Path(tempfile.mkdtemp())
        model = tmp / "modelo.xlsx"
        sheet.write_model(model)
        book = openpyxl.load_workbook(model)
        clubs = {row[6].value: row for row in book["Clubes"].iter_rows(min_row=2)}
        self.assertIn("River Plate (URU)", clubs)                   # o mesmo nome em dois países
        clubs["River Plate (ARG)"][7].value = "Marcelo Gallardo"
        book["Jogadores"].append(["River Plate (ARG)", 1, "Franco Armani", None, None, "1986-10-16", None, "GK"])
        book["Jogadores"].append(["Peñarol", None, "Leo Fernández", None, None, "1998-11-08", None, "AM RW"])
        book["Jogadores"].append(["River Plate", None, "Ambíguo", None, None, "2000-01-01", None, "CM"])
        book["Carreiras"].append(["River Plate (ARG)", "Franco Armani", "1986-10-16", "2018-", "River Plate"])
        book.save(model)
        paths = {"DATA": tmp, "IDENTITIES": tmp / "i.csv", "PLAYERS": tmp / "j.xlsx", "COACHES": tmp / "t.csv",
                 "CAREERS": tmp / "c.csv"}
        with mock.patch.multiple(ex, **paths):
            errors = sheet.build(model)
            players = list(openpyxl.load_workbook(tmp / "j.xlsx").active.iter_rows(values_only=True))
            rows = {r["clube_2026"]: r for r in ex.identity_rows()}
        self.assertEqual(len(errors), 1)                            # "River Plate" sem o país
        self.assertEqual([(p[0], p[2], p[6]) for p in players[1:]],
                         [("River Plate (ARG)", "Franco Armani", "Argentina"), ("Peñarol (URU)", "Leo Fernández", "Uruguai")])
        self.assertEqual(rows["River Plate (ARG)"]["uid_jogo"], "94")
        self.assertIn("Marcelo Gallardo", (tmp / "t.csv").read_text(encoding="utf-8-sig"))
        self.assertIn("2018–,River Plate", (tmp / "c.csv").read_text(encoding="utf-8-sig"))


class CareerClubTest(unittest.TestCase):
    def test_2026_name_of_the_club(self):
        import estrangeiros_planilha as sheet

        def club(full, country, long, short):
            return full, country, ex._keys({"long": long, "short": short}) | {hj._key(full.split(" (")[0])}
        known = [club("River Plate (ARG)", "AR", "River Plate", "River"),
                 club("River Plate (URU)", "UY", "CA River Plate", "River Plate (URU)"),
                 club("Nacional (URU)", "UY", "C Nacional F", "Nacional"),
                 club("Nacional (PAR)", "PY", "C Nacional FC", "Nacional (PAR)"),
                 club("Cerro (URU)", "UY", "CA Cerro", "Cerro"),
                 club("Cerro Porteño (PAR)", "PY", "C Cerro Porteño", "Cerro Porteño"),
                 club("Vélez Sarsfield (ARG)", "AR", "Vélez Sarsfield", "Vélez"),
                 ("Internacional", "BR", {hj._key("Internacional")}),
                 ("Inter de Limeira", "BR", {hj._key("Inter de Limeira"), hj._key("AA Internacional")})]
        find = sheet.career_club
        self.assertEqual(find("CA River Plate", "AR", known), "River Plate (ARG)")
        self.assertEqual(find("Club Nacional", "PY", known), "Nacional (PAR)")      # o do país do jogador
        self.assertEqual(find("Club Nacional Asunción", "PY", known), "Nacional (PAR)")
        self.assertEqual(find("Club Cerro Porteño Asunción", "PY", known), "Cerro Porteño (PAR)")
        self.assertEqual(find("CA Vélez Sarsfield II", "AR", known), "Vélez Sarsfield (ARG) II")
        self.assertEqual(find("Sport Club Internacional", "AR", known), "Internacional")
        self.assertEqual(find("Real Madrid", "AR", known), "Real Madrid")
